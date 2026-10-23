"""Превращает итог разговора из простого markdown в аккуратный .docx.

Запуск:
    python make_docx.py "results/Голос 002/итог.md" ["results/Голос 002/итог.docx"]

Поддерживается: # заголовок, ## / ### подзаголовки, списки «- » и «1. », цитаты «> »,
таблицы «| a | b |», **жирный** внутри строки, обычные абзацы.

Сайт вызывает build_docx(): итог из markdown и, по желанию, полная расшифровка в конце.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x5F, 0x6B)


def add_runs(par, text: str, italic=False, color=None):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = par.add_run(part)
        run.bold = i % 2 == 1
        run.italic = italic
        if color:
            run.font.color.rgb = color


def shade(cell, hex_fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # не даём теме перебить шрифт
        st.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(6)


def flush_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"\|?\s*:?-{2,}.*", r.replace(" ", ""))]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    ncols = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(cells):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            par = cell.paragraphs[0]
            add_runs(par, row[ci] if ci < len(row) else "")
            if ri == 0:
                for run in par.runs:
                    run.bold = True
                shade(cell, "E8EEF5")
    doc.add_paragraph()


TRANSCRIPT_LINE = re.compile(r"^\[(\d{1,2}:\d{2}(?::\d{2})?)\]\s*([^:]{1,60}):\s*(.*)$")


def add_transcript(doc, transcript: str, title: str):
    """Расшифровка с новой страницы: время серым, говорящий жирным, дальше текст."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(title, level=1)
    for raw in transcript.splitlines():
        line = raw.strip()
        if not line:
            continue
        par = doc.add_paragraph()
        m = TRANSCRIPT_LINE.match(line)
        if not m:
            run = par.add_run(line)
            if line.startswith("---"):
                run.italic = True
                run.font.color.rgb = MUTED
            continue
        stamp = par.add_run(f"[{m[1]}] ")
        stamp.font.color.rgb = MUTED
        stamp.font.size = Pt(9.5)
        par.add_run(f"{m[2].strip()}: ").bold = True
        par.add_run(m[3])


def build_docx(md_text: str, out_path: Path, transcript: str | None = None,
               transcript_title: str = "Полная расшифровка"):
    doc = Document()
    setup(doc)
    table_rows = []

    for raw in md_text.splitlines():
        line = raw.rstrip()
        if line.startswith("|"):
            table_rows.append(line)
            continue
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:].strip(), style="Title")
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\s*[-*] ", line):
            level = 2 if line.startswith("  ") else 1
            par = doc.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
            add_runs(par, re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\s*\d+[.)] ", line):
            par = doc.add_paragraph(style="List Number")
            add_runs(par, re.sub(r"^\s*\d+[.)] ", "", line))
        elif line.startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.8)
            add_runs(par, line.lstrip("> ").strip(), italic=True, color=MUTED)
        elif line.startswith("_") and line.endswith("_"):
            par = doc.add_paragraph()
            add_runs(par, line.strip("_"), italic=True, color=MUTED)
        else:
            add_runs(doc.add_paragraph(), line)

    if table_rows:
        flush_table(doc, table_rows)
    if transcript:
        add_transcript(doc, transcript, transcript_title)
    doc.save(out_path)


def convert(md_path: Path, out_path: Path):
    build_docx(md_path.read_text(encoding="utf-8-sig"), out_path)
    print(f"Готово: {out_path}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    convert(src, Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx"))
